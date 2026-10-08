#!/usr/bin/env python3
"""Build an editable PowerPoint reconstruction of the ForeTac Overview film."""

from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement

W = Inches(13.333333)
H = Inches(7.5)
NAVY = RGBColor(15, 23, 42)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(226, 232, 240)
ORANGE = RGBColor(234, 88, 12)
PURPLE = RGBColor(124, 58, 237)


def add_text(slide, text, x, y, w, h, size, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, font="DejaVu Sans"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title_overlay(slide, title, subtitle=None):
    panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(10.8), Inches(1.95)
    )
    panel.name = "Editable title panel - 28 percent transparent"
    panel.fill.solid()
    panel.fill.fore_color.rgb = NAVY
    # python-pptx has no fill transparency property; write actual DrawingML alpha.
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "72000")
    panel.element.xpath("./p:spPr/a:solidFill/a:srgbClr")[0].append(alpha)
    panel.line.fill.background()
    add_text(slide, title, 0.78, 0.63, 10.25, 0.67, 36, WHITE, True)
    if subtitle:
        add_text(slide, subtitle, 0.8, 1.42, 10.15, 0.38, 16, LIGHT)


def add_media(slide, path, poster, x=0, y=0, w=13.333333, h=7.5):
    movie = slide.shapes.add_movie(
        str(path), Inches(x), Inches(y), Inches(w), Inches(h),
        poster_frame_image=str(poster), mime_type="video/mp4",
    )
    movie.name = path.name
    for condition in slide.element.xpath(".//p:video/p:cMediaNode/p:cTn/p:stCondLst/p:cond"):
        condition.set("delay", "0")


def add_image(slide, path, x=0.5, y=1.5, w=12.333333, h=5.55):
    # Convert web assets to PNG and preserve aspect ratio within the content area.
    with Image.open(path) as source:
        img = source.convert("RGB")
        img.thumbnail((2400, 1800))
        factor = min(w / img.width, h / img.height)
        picture_w, picture_h = img.width * factor, img.height * factor
        stream = BytesIO()
        img.save(stream, format="PNG")
    stream.seek(0)
    shape = slide.shapes.add_picture(
        stream, Inches(x + (w - picture_w) / 2), Inches(y + (h - picture_h) / 2),
        Inches(picture_w), Inches(picture_h),
    )
    shape.name = path.name


def add_heading(slide, title, subtitle):
    add_text(slide, title, 0.5, 0.62, 12.25, 0.4, 23, NAVY, True)
    add_text(slide, subtitle, 0.5, 1.05, 12.25, 0.28, 13, RGBColor(71, 85, 105))


def add_clean_slide(prs, media, images, title, subtitle, image=None,
                    video=None, poster=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if video:
        add_media(slide, media / video,
                  media / (poster or "hero_montage_no_title_poster.jpg"))
    else:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
        add_image(slide, images / image)
    add_heading(slide, title, subtitle)
    return slide


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--editable-cover", action="store_true",
                        help="Add a native editable title panel to the clean opening video.")
    args = parser.parse_args()

    repo = args.repo
    media = repo / "static/videos/overview"
    images = repo / "static/images/manuscript"
    required = [
        media / "hero_montage_no_title.mp4",
        media / "hero_montage_no_title_poster.jpg",
        images / "foretac_teaser_fig1.webp",
        images / "foretac_method_fig2.webp",
        images / "foretac_tasks_fig3.webp",
        images / "foretac_guidance_fig4_r12.png",
        media / "segments/05_foresight_clean.mp4",
        media / "segments/05_foresight_clean_poster.jpg",
    ]
    missing = [path for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing media: " + ", ".join(map(str, missing)))

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # Clean opening by default; the alternate cover uses native editable objects.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_media(slide, media / "hero_montage_no_title.mp4",
              media / "hero_montage_no_title_poster.jpg")
    if args.editable_cover:
        add_title_overlay(slide, "ForeTac",
                          "Predictive Contact Guidance for Generative Robot Policies")
        add_text(slide, "Four real-robot contact regimes", 0.8, 1.94, 10.1, 0.28, 13, LIGHT)

    add_clean_slide(
        prs, media, images, "Predict the contact before it happens",
        "Anticipate the tactile consequence of a candidate action",
        image="foretac_teaser_fig1.webp")
    add_clean_slide(
        prs, media, images, "Represent - predict - evaluate - guide",
        "Frozen models guide the evolving generative sampling state",
        image="foretac_method_fig2.webp")
    add_clean_slide(
        prs, media, images, "Four real-robot contact regimes",
        "Board wiping - vase wiping - card swiping - fragile chip grasping",
        image="foretac_tasks_fig3.webp")
    add_clean_slide(
        prs, media, images, "Action-conditioned tactile foresight",
        "Ground-truth future - predicted future - error (qualitative, H=16)",
        video="segments/05_foresight_clean.mp4",
        poster="segments/05_foresight_clean_poster.jpg")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
    add_image(slide, images / "foretac_guidance_fig4_r12.png")
    add_heading(slide, "Offline guidance diagnostics",
                "Scorer-margin changes and scaled quality gradients")

    # Metrics are native editable text rather than burned-in pixels.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_text(slide, "ForeTac on four real-robot tasks", 0.85, 0.7, 11.7, 0.6, 34, WHITE, True)
    add_text(slide, "DP 36.3 / 20.0  →  DP + ForeTac 71.3 / 66.3   (SR / CSR)",
             1.0, 1.65, 11.2, 0.35, 17, LIGHT)
    add_text(slide, "20 trials per task; averages over four tasks",
             1.0, 2.06, 11.2, 0.28, 13, LIGHT)
    add_text(slide, "Average task success rate (SR)", 1.0, 2.55, 4.6, 0.35, 16, LIGHT)
    add_text(slide, "71.3%", 1.0, 2.95, 4.6, 0.8, 42, ORANGE, True)
    add_text(slide, "Average contact success rate (CSR)", 7.0, 2.55, 5.2, 0.35, 16, LIGHT)
    add_text(slide, "66.3%", 7.0, 2.95, 4.6, 0.8, 42, PURPLE, True)
    add_text(slide, "vs RDP  +12.5 / +17.5 pp (SR / CSR)",
             1.0, 4.55, 10.5, 0.4, 20, WHITE)
    add_text(slide, "pi0.5 46.3 / 36.3  →  + ForeTac 67.5 / 62.5   (SR / CSR)",
             1.0, 5.25, 11.0, 0.4, 17, LIGHT)
    add_text(slide, "Diffusion Policy and flow-matching backbones",
             1.0, 5.76, 11.0, 0.28, 13, LIGHT)
    add_text(slide, "Predict contact. Guide action. Execute and replan.",
             1.0, 6.35, 11.0, 0.45, 21, WHITE)

    prs.core_properties.title = "ForeTac Overview (Editable)"
    prs.core_properties.subject = "Editable reconstruction of the ForeTac Overview film"
    prs.core_properties.author = "ForeTac"
    for page, time_span in zip(prs.slides, ("0-5", "5-11", "11-19", "19-25",
                                          "25-32", "32-38", "38-45")):
        page.notes_slide.notes_text_frame.text = (
            f"Overview film interval: {time_span} seconds. "
            "Headings and metrics are editable text. Figures remain image assets; "
            "movie content remains embedded video. This deck is not a single embedded film."
        )
    prs.core_properties.comments = (
        "The opening montage and foresight video are embedded; overlay panels and "
        "text are native editable PowerPoint objects."
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
